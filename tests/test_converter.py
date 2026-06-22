"""Tests for the core conversion pipeline."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from html_to_pptx import convert, extract_measurements, render_pptx


# ---------------------------------------------------------------------------
# End-to-end (these launch a browser — kept to a minimum)
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_convert_creates_valid_pptx(minimal_html: Path, tmp_pptx: Path):
    result = await convert(str(minimal_html), str(tmp_pptx))
    assert result == str(tmp_pptx)
    assert Path(result).stat().st_size > 0

    prs = Presentation(str(tmp_pptx))
    assert len(prs.slides) == 2


@pytest.mark.asyncio
async def test_file_not_found_raises():
    with pytest.raises(FileNotFoundError):
        await extract_measurements("/nonexistent/path.html")


@pytest.mark.asyncio
async def test_empty_html_raises(tmp_path: Path):
    empty = tmp_path / "empty.html"
    empty.write_text("<html><body></body></html>")
    with pytest.raises(ValueError, match="No slides found"):
        await extract_measurements(str(empty))


# ---------------------------------------------------------------------------
# Measurement tests (reuse session-scoped fixture — no browser per test)
# ---------------------------------------------------------------------------


def test_measurements_returns_correct_slide_count(minimal_measurements: list[dict]):
    assert len(minimal_measurements) == 2


def test_measurements_have_required_keys(minimal_measurements: list[dict]):
    for slide in minimal_measurements:
        assert "index" in slide
        assert "elements" in slide
        assert "backgroundColor" in slide


def test_measurements_capture_text(minimal_measurements: list[dict]):
    def find_text(elements: list[dict], target: str) -> bool:
        for el in elements:
            if target in el.get("text", ""):
                return True
            if find_text(el.get("children", []), target):
                return True
        return False

    assert find_text(minimal_measurements[0]["elements"], "Title Slide")
    assert find_text(minimal_measurements[1]["elements"], "Content Slide")


def test_measurements_capture_background_color(minimal_measurements: list[dict]):
    assert minimal_measurements[0]["backgroundColor"] is not None


# ---------------------------------------------------------------------------
# Rendering tests (pure python-pptx, no browser)
# ---------------------------------------------------------------------------


def test_render_creates_correct_slide_count(minimal_measurements: list[dict]):
    prs = render_pptx(minimal_measurements)
    assert len(prs.slides) == 2


def test_render_sets_slide_dimensions(minimal_measurements: list[dict]):
    prs = render_pptx(minimal_measurements)
    assert abs(prs.slide_width - Inches(13.333)) < Inches(0.01)
    assert prs.slide_height == Inches(7.5)


def test_render_adds_shapes(minimal_measurements: list[dict]):
    prs = render_pptx(minimal_measurements)
    for slide in prs.slides:
        assert len(slide.shapes) > 0


def test_render_output_is_saveable(minimal_measurements: list[dict], tmp_pptx: Path):
    prs = render_pptx(minimal_measurements)
    prs.save(str(tmp_pptx))
    reloaded = Presentation(str(tmp_pptx))
    assert len(reloaded.slides) == 2
