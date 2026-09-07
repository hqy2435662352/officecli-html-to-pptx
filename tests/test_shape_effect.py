"""Regression tests for https://github.com/Design-Arena/html-to-pptx/issues/2.

Autoshapes must not inherit the Office theme's default shadow effect
(``<p:style><a:effectRef idx="2">``) — the input HTML never authors a shadow.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from html_to_pptx import convert
from html_to_pptx import converter as C


def test_strip_theme_effect_clears_inherited_shadow():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1),
    )
    # python-pptx's add_shape injects <p:style><a:effectRef idx="2"> -> shadow inherited
    assert shape.shadow.inherit is True
    C._strip_theme_effect(shape)
    assert shape.shadow.inherit is False
    # an explicit empty effect list overrides the theme effectRef in spPr
    effect_lst = shape._element.spPr.find(qn("a:effectLst"))
    assert effect_lst is not None
    assert len(effect_lst) == 0


@pytest.mark.asyncio
async def test_converted_autoshapes_have_no_inherited_theme_effect(tmp_path: Path):
    html = tmp_path / "card.html"
    html.write_text(
        """<!DOCTYPE html><html><head><style>
        * { margin:0; padding:0; box-sizing:border-box; }
        body { overflow:hidden; }
        .slide { width:1920px; height:1080px; display:none; flex-direction:column; }
        .slide.active { display:flex; }
        .card { width:400px; height:200px; border:1px solid #d8d8d8;
                border-radius:18px; background:#ffffff; margin:100px; }
        </style></head><body>
        <section class="slide active"><div class="card">hello</div></section>
        </body></html>""",
        encoding="utf-8",
    )
    pptx_path = tmp_path / "out.pptx"
    await convert(str(html), str(pptx_path))

    prs = Presentation(str(pptx_path))
    autoshapes = [
        shape
        for slide in prs.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
    ]
    assert autoshapes, "expected at least one autoshape in the output"
    for shape in autoshapes:
        effect_lst = shape._element.spPr.find(qn("a:effectLst"))
        assert effect_lst is not None, (
            f"autoshape id={shape.shape_id} still inherits the theme effect"
        )